import openpyxl

from core.exporter import ALL_PARCELS_SHEET_NAME, SUMMARY_SHEET_NAME, export
from core.merger import UNKNOWN_NATIONAL_ID, apply_manual_national_ids, build_manual_review_people
from core.models import AssociationInfo, AssociationType, BasinInfo, MergeResult, Parcel


def _make_parcel(holding, holder, basin, feddan=0.0, qirat=0.0, sahm=0.0, border_n=".") -> Parcel:
    return Parcel(
        directorate="الدقهليه",
        administration="اجا",
        association_name="شنشا-الائتمان الزراعي",
        association_type="credit",
        association_code="323925",
        basin_name=basin,
        basin_code=f"code-{basin}",
        holding_number=holding,
        unified_holding_id="06-3230-00323925-000001",
        registry_page="1",
        national_id="12345678901234",
        holder_name=holder,
        parcel_count_in_holding=1,
        land_number="10000000",
        area_feddan=feddan,
        area_qirat=qirat,
        area_sahm=sahm,
        area_m2=0.0,
        border_north=border_n,
        border_west=".",
        border_south=".",
        border_east=".",
    )


def _make_result() -> MergeResult:
    association = AssociationInfo(
        name="شنشا-الائتمان الزراعي",
        type=AssociationType.CREDIT,
        code="323925",
        directorate="الدقهليه",
        administration="اجا",
    )
    basins = [BasinInfo(name="الدماسه", code="code-الدماسه")]
    parcels = [
        _make_parcel("1", "احمد محمد", "الدماسه", feddan=1.0),
        _make_parcel("2", "محمد احمد", "الدماسه", feddan=2.0, border_n="احمد محمد"),
    ]
    return MergeResult(parcels=parcels, association=association, basins=basins, unmatched_count=0)


def test_export_creates_expected_sheets(tmp_path):
    result = _make_result()
    output_path = tmp_path / "output.xlsx"

    export(result, output_path)

    wb = openpyxl.load_workbook(output_path)
    assert ALL_PARCELS_SHEET_NAME in wb.sheetnames
    assert SUMMARY_SHEET_NAME in wb.sheetnames
    assert "الدماسه" in wb.sheetnames


def test_all_parcels_sheet_is_rtl_and_has_all_rows(tmp_path):
    result = _make_result()
    output_path = tmp_path / "output.xlsx"

    export(result, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb[ALL_PARCELS_SHEET_NAME]
    assert ws.sheet_view.rightToLeft is True
    # header + separator + 2 parcel rows
    assert ws.max_row == 4


def test_border_hyperlink_resolved(tmp_path):
    result = _make_result()
    output_path = tmp_path / "output.xlsx"

    export(result, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb[ALL_PARCELS_SHEET_NAME]
    # second parcel row's north border references the first parcel by name
    linked_cell = ws.cell(row=4, column=19)
    assert linked_cell.hyperlink is not None


def test_manual_national_id_is_written_without_changing_other_parcels(tmp_path):
    result = _make_result()
    result.parcels[0].national_id = UNKNOWN_NATIONAL_ID
    person = build_manual_review_people(result.parcels)[0]
    apply_manual_national_ids(result, {person.key: "29510251202211"})
    output_path = tmp_path / "output.xlsx"

    export(result, output_path)

    ws = openpyxl.load_workbook(output_path)[ALL_PARCELS_SHEET_NAME]
    assert ws.cell(row=3, column=11).value == "29510251202211"
    assert ws.cell(row=4, column=11).value == "12345678901234"


def test_export_removes_only_rows_identical_in_every_exported_field(tmp_path):
    result = _make_result()
    duplicate = _make_parcel("1", "احمد محمد", "الدماسه", feddan=1.0)
    result.parcels.append(duplicate)

    removed = export(result, tmp_path / "output.xlsx")

    assert removed == 1
    assert result.duplicates_removed == 1
    assert len(result.parcels) == 2


def test_export_keeps_rows_when_one_exported_field_differs(tmp_path):
    result = _make_result()
    similar = _make_parcel("1", "احمد محمد", "الدماسه", feddan=1.0)
    similar.border_east = "حد مختلف"
    result.parcels.append(similar)

    removed = export(result, tmp_path / "output.xlsx")

    assert removed == 0
    assert len(result.parcels) == 3
