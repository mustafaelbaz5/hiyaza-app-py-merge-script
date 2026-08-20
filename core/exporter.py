"""Writes the final merged Excel report: all-parcels sheet, per-basin sheets,
and a basin summary sheet, with border-cell hyperlinks."""

import logging
from pathlib import Path

from openpyxl import Workbook

from core import export_sheet as sheet
from core import export_styles as styles
from core.exceptions import MergerError
from core.hyperlinks import resolve_border_hyperlink
from core.models import MergeResult, Parcel

logger = logging.getLogger(__name__)

ALL_PARCELS_SHEET_NAME = "كل الحيازات"
SUMMARY_SHEET_NAME = "ملخص الأحواض"
_MAX_SHEET_NAME_LEN = 31


def export(result: MergeResult, output_path: Path) -> None:
    try:
        parcels_by_basin = _group_by_basin_ordered(result)
        wb = Workbook()
        wb.remove(wb.active)

        row_to_excel = _write_all_parcels_sheet(wb, parcels_by_basin)
        for basin_name, parcels in parcels_by_basin.items():
            _write_basin_sheet(wb, basin_name, parcels)
        _write_summary_sheet(wb, parcels_by_basin)

        _apply_border_hyperlinks(wb[ALL_PARCELS_SHEET_NAME], result.parcels, row_to_excel)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
    except MergerError:
        raise
    except Exception as e:
        logger.exception("Unexpected error writing Excel output")
        raise MergerError("حدث خطأ غير متوقع أثناء كتابة ملف Excel") from e


def _holding_sort_key(parcel: Parcel) -> int:
    try:
        return int(parcel.holding_number)
    except (TypeError, ValueError):
        return 0


def _group_by_basin_ordered(result: MergeResult) -> dict[str, list[Parcel]]:
    """Groups parcels by basin (ordered by basin code) then sorts each
    group's parcels by holding number, per the required export order."""
    grouped: dict[str, list[Parcel]] = {}
    for parcel in result.parcels:
        grouped.setdefault(parcel.basin_name, []).append(parcel)

    basin_order = [b.name for b in result.basins if b.name in grouped]
    remaining = [name for name in grouped if name not in basin_order]
    ordered_names = basin_order + remaining

    return {
        name: sorted(grouped[name], key=_holding_sort_key) for name in ordered_names
    }


def _sheet_safe_name(name: str) -> str:
    invalid = set('[]:*?/\\')
    cleaned = "".join(c for c in name if c not in invalid)
    return cleaned[:_MAX_SHEET_NAME_LEN]


def _write_all_parcels_sheet(
    wb: Workbook, parcels_by_basin: dict[str, list[Parcel]]
) -> dict[int, int]:
    ws = wb.create_sheet(ALL_PARCELS_SHEET_NAME)
    sheet.setup_sheet(ws)
    sheet.write_header(ws)

    row_to_excel: dict[int, int] = {}
    excel_row = 2
    for basin_name, parcels in parcels_by_basin.items():
        sheet.write_separator_row(ws, excel_row, f"حوض {basin_name}")
        excel_row += 1
        colors = styles.BASIN_ROW_COLORS.get(basin_name, styles.DEFAULT_ROW_COLORS)
        for i, parcel in enumerate(parcels):
            color = colors[i % 2]
            sheet.write_parcel_row(ws, excel_row, parcel, color)
            row_to_excel[id(parcel)] = excel_row
            excel_row += 1
    return row_to_excel


def _write_basin_sheet(wb: Workbook, basin_name: str, parcels: list[Parcel]) -> None:
    ws = wb.create_sheet(_sheet_safe_name(basin_name))
    sheet.setup_sheet(ws)
    sheet.write_header(ws)
    colors = styles.BASIN_ROW_COLORS.get(basin_name, styles.DEFAULT_ROW_COLORS)
    for i, parcel in enumerate(parcels):
        sheet.write_parcel_row(ws, i + 2, parcel, colors[i % 2])


def _write_summary_sheet(wb: Workbook, parcels_by_basin: dict[str, list[Parcel]]) -> None:
    ws = wb.create_sheet(SUMMARY_SHEET_NAME)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    headers = ["اسم الحوض", "كود الحوض", "عدد القطع", "إجمالي المساحة م²"]
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = styles.header_font("header_basin")
        cell.fill = styles.header_fill("header_basin")
        cell.alignment = styles.CENTER_ALIGN
    ws.auto_filter.ref = "A1:D1"

    for row_idx, (basin_name, parcels) in enumerate(parcels_by_basin.items(), start=2):
        total_m2 = sum(p.area_m2 for p in parcels)
        basin_code = parcels[0].basin_code if parcels else ""
        ws.cell(row=row_idx, column=1, value=basin_name).font = styles.body_font()
        ws.cell(row=row_idx, column=2, value=basin_code).font = styles.body_font()
        ws.cell(row=row_idx, column=3, value=len(parcels)).font = styles.body_font()
        m2_cell = ws.cell(row=row_idx, column=4, value=round(total_m2, 2))
        m2_cell.font = styles.body_font()
        m2_cell.number_format = styles.NUMBER_FORMAT_M2


def _apply_border_hyperlinks(
    ws, parcels: list[Parcel], row_to_excel: dict[int, int]
) -> None:
    name_index: dict[str, list[tuple[int, str]]] = {}
    for parcel in parcels:
        name_index.setdefault(parcel.holder_name, []).append(
            (id(parcel), parcel.basin_name)
        )

    border_columns = {
        "border_north": 19,
        "border_west": 20,
        "border_south": 21,
        "border_east": 22,
    }
    for parcel in parcels:
        excel_row = row_to_excel.get(id(parcel))
        if excel_row is None:
            continue
        for key, col_idx in border_columns.items():
            text = getattr(parcel, key)
            ref, is_multi = resolve_border_hyperlink(
                text, name_index, parcel.basin_name, row_to_excel
            )
            if ref is None:
                continue
            cell = ws.cell(row=excel_row, column=col_idx)
            cell.hyperlink = ref
            cell.fill = styles.multi_match_fill() if is_multi else styles.link_fill()
