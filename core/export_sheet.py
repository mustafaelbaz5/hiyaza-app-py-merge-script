"""Writes parcel rows into a single worksheet with headers and formatting."""

from openpyxl.worksheet.worksheet import Worksheet

from core import export_styles as styles
from core.models import Parcel

_ASSOC_TYPE_LABELS = {"credit": "ائتمان زراعي", "reform": "إصلاح زراعي"}


def setup_sheet(ws: Worksheet) -> None:
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = styles.HEADER_ROW_HEIGHT


def write_header(ws: Worksheet) -> None:
    for col_idx, (_key, label, color_group) in enumerate(styles.COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = styles.header_font(color_group)
        cell.fill = styles.header_fill(color_group)
        cell.alignment = styles.CENTER_ALIGN
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(styles.COLUMNS)).coordinate}"


def parcel_value(parcel: Parcel, key: str):
    if key == "association_type_label":
        return _ASSOC_TYPE_LABELS.get(parcel.association_type, parcel.association_type)
    if key in ("area_feddan", "area_qirat", "area_sahm", "area_m2"):
        value = getattr(parcel, key)
        return None if value == 0 else value
    return getattr(parcel, key)


def write_parcel_row(ws: Worksheet, excel_row: int, parcel: Parcel, row_color: str) -> None:
    ws.row_dimensions[excel_row].height = styles.DATA_ROW_HEIGHT
    fill = styles.row_fill(row_color)
    for col_idx, (key, _label, _group) in enumerate(styles.COLUMNS, start=1):
        cell = ws.cell(row=excel_row, column=col_idx, value=parcel_value(parcel, key))
        cell.font = styles.body_font()
        cell.alignment = styles.RIGHT_ALIGN
        cell.fill = fill
        _apply_column_format(cell, key)


def _apply_column_format(cell, key: str) -> None:
    if key in ("area_feddan", "area_qirat", "area_sahm"):
        cell.number_format = styles.NUMBER_FORMAT_QUANTITY
    elif key == "area_m2":
        cell.number_format = styles.NUMBER_FORMAT_M2
        cell.fill = styles.m2_fill()


def write_separator_row(ws: Worksheet, excel_row: int, label: str) -> None:
    ws.row_dimensions[excel_row].height = styles.DATA_ROW_HEIGHT
    ws.merge_cells(
        start_row=excel_row,
        start_column=1,
        end_row=excel_row,
        end_column=len(styles.COLUMNS),
    )
    cell = ws.cell(row=excel_row, column=1, value=label)
    cell.font = styles.separator_font()
    cell.fill = styles.separator_fill()
    cell.alignment = styles.CENTER_ALIGN
