"""Color palette, fonts, and column layout for the Excel export."""

from openpyxl.styles import Alignment, Font, PatternFill

FONT_NAME = "Arial"

COLORS = {
    "header_identity": "1B5E20",
    "header_basin": "004D40",
    "header_holding": "0D47A1",
    "header_parcel": "4A148C",
    "header_borders": "37474F",
    "m2_cell": "FFF9C4",
    "link_cell": "E3F2FD",
    "multi_match": "FFF9C4",
    "separator": "263238",
    "row_alt_a": "FFFFFF",
    "row_alt_b": "F5F5F5",
}

# (column key, header text, header color group)
COLUMNS: list[tuple[str, str, str]] = [
    ("directorate", "المديرية", "header_identity"),
    ("administration", "الإدارة", "header_identity"),
    ("association_name", "الجمعية", "header_identity"),
    ("association_type_label", "نوع الجمعية", "header_identity"),
    ("association_code", "كود الجمعية", "header_identity"),
    ("basin_name", "اسم الحوض", "header_basin"),
    ("basin_code", "كود الحوض", "header_basin"),
    ("holding_number", "رقم الحيازة", "header_holding"),
    ("unified_holding_id", "الرقم الموحد", "header_holding"),
    ("registry_page", "رقم الصفحة", "header_holding"),
    ("national_id", "الرقم القومي", "header_holding"),
    ("holder_name", "اسم الحائز", "header_holding"),
    ("parcel_count_in_holding", "عدد القطع", "header_holding"),
    ("land_number", "رقم الأرض", "header_parcel"),
    ("area_feddan", "فدان", "header_parcel"),
    ("area_qirat", "قيراط", "header_parcel"),
    ("area_sahm", "سهم", "header_parcel"),
    ("area_m2", "مساحة القطعة - م²", "header_parcel"),
    ("border_north", "حد بحري", "header_borders"),
    ("border_west", "حد غربي", "header_borders"),
    ("border_south", "حد قبلي", "header_borders"),
    ("border_east", "حد شرقي", "header_borders"),
]

BORDER_COLUMN_KEYS = {"border_north", "border_west", "border_south", "border_east"}

COLUMN_WIDTHS: dict[str, float] = {
    "directorate": 14,
    "administration": 12,
    "association_name": 26,
    "association_type_label": 14,
    "association_code": 12,
    "basin_name": 16,
    "basin_code": 22,
    "holding_number": 12,
    "unified_holding_id": 26,
    "registry_page": 10,
    "national_id": 16,
    "holder_name": 26,
    "parcel_count_in_holding": 10,
    "land_number": 12,
    "area_feddan": 10,
    "area_qirat": 10,
    "area_sahm": 10,
    "area_m2": 18,
    "border_north": 20,
    "border_west": 20,
    "border_south": 20,
    "border_east": 20,
}
NUMBER_FORMAT_QUANTITY = "#,##0.###"
NUMBER_FORMAT_M2 = '#,##0.00 "م²"'

HEADER_ROW_HEIGHT = 34
DATA_ROW_HEIGHT = 15

BASIN_ROW_COLORS: dict[str, tuple[str, str]] = {
    "الباشا": ("E8F5E9", "F1F8E9"),
    "البحيره": ("E3F2FD", "EAF4FD"),
    "البراوي": ("FFF3E0", "FFF8E6"),
    "الخلف": ("F3E5F5", "F8EAFB"),
    "الدماسه": ("E0F2F1", "E8F6F4"),
    "العدل": ("FFFDE7", "FFFFF3"),
    "العمده": ("EFEBE9", "F5F1F0"),
    "الفوقانيه": ("E8EAF6", "EEF0FA"),
    "القصير": ("FCE4EC", "FDEDF2"),
    "القطع": ("F1F8E9", "F6FBF0"),
    "داير الناحيه": ("ECEFF1", "F4F6F7"),
    "عبيد اول": ("FFF8E1", "FFFBEE"),
    "عبيد ثاني": ("E1F5FE", "EAFAFF"),
}
DEFAULT_ROW_COLORS = ("FFFFFF", "F5F5F5")


def header_font(color_key: str) -> Font:
    return Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)


def header_fill(color_key: str) -> PatternFill:
    return PatternFill(
        start_color=COLORS[color_key], end_color=COLORS[color_key], fill_type="solid"
    )


def body_font() -> Font:
    return Font(name=FONT_NAME, size=10)


def separator_font() -> Font:
    return Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)


def separator_fill() -> PatternFill:
    return PatternFill(
        start_color=COLORS["separator"], end_color=COLORS["separator"], fill_type="solid"
    )


def m2_fill() -> PatternFill:
    return PatternFill(
        start_color=COLORS["m2_cell"], end_color=COLORS["m2_cell"], fill_type="solid"
    )


def link_fill() -> PatternFill:
    return PatternFill(
        start_color=COLORS["link_cell"], end_color=COLORS["link_cell"], fill_type="solid"
    )


def multi_match_fill() -> PatternFill:
    return PatternFill(
        start_color=COLORS["multi_match"],
        end_color=COLORS["multi_match"],
        fill_type="solid",
    )


def row_fill(color_hex: str) -> PatternFill:
    return PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")


RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
